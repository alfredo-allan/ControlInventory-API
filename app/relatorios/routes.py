# Relatórios em .docx — pedido do usuário ("os relatórios que devemos
# gerar é interessante que sejam em docx"). Gerados aqui no backend (não no
# navegador) porque é aqui que mora o dado de verdade, cross-dispositivo —
# um relatório gerado no frontend só veria o que aquele aparelho sincronizou.
# `python-docx` monta o .docx inteiro em memória (BytesIO), devolvido como
# download — nada fica salvo em disco no servidor.

from datetime import datetime, timezone
from io import BytesIO

from flask import Blueprint, g, jsonify, request, send_file
from docx import Document
from docx.shared import Pt

from app.models import Avaria, Mercado, PedidoItem, Produto
from app.auth.utils import login_obrigatorio
from app.metricas.routes import calcular_perdas_por_mercado

bp = Blueprint("relatorios", __name__, url_prefix="/api/relatorios")


def _cabecalho(doc: Document, titulo: str, subtitulo: str) -> None:
    doc.add_heading(titulo, level=1)
    p = doc.add_paragraph(subtitulo)
    p.runs[0].italic = True
    agora = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    rodape = doc.add_paragraph(f"Gerado em {agora}")
    rodape.runs[0].font.size = Pt(9)
    doc.add_paragraph()


def _tabela(doc: Document, colunas: list[str]) -> "object":
    tabela = doc.add_table(rows=1, cols=len(colunas))
    tabela.style = "Light Grid Accent 1"
    cabecalho = tabela.rows[0].cells
    for i, titulo in enumerate(colunas):
        cabecalho[i].text = titulo
        for run in cabecalho[i].paragraphs[0].runs:
            run.bold = True
    return tabela


def _resposta_docx(doc: Document, nome_arquivo: str):
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=nome_arquivo,
    )


@bp.get("/avarias.docx")
@login_obrigatorio()
def relatorio_avarias():
    """Aberto pros dois perfis — mesmo nível de sensibilidade da listagem
    crua (GET /api/avarias): um operador pode querer o próprio relatório
    como comprovante do que registrou."""
    avarias = Avaria.query.filter_by(empresa_id=g.empresa_id).order_by(Avaria.registrado_em.desc()).all()

    doc = Document()
    _cabecalho(doc, "Relatório de Avarias e Quebras", f"{len(avarias)} registro(s) — histórico completo")

    if not avarias:
        doc.add_paragraph("Nenhuma avaria registrada até o momento.")
    else:
        tabela = _tabela(doc, ["Data", "Mercado", "EAN", "Produto", "Qtd.", "Operador", "Motivo"])
        for a in avarias:
            linha = tabela.add_row().cells
            linha[0].text = a.registrado_em.strftime("%d/%m/%Y %H:%M") if a.registrado_em else "—"
            linha[1].text = a.nome_cliente or a.mercado_slug
            linha[2].text = a.codigo_ean or "—"
            linha[3].text = a.descricao or "—"
            linha[4].text = f"{a.quantidade} {'cx' if a.quantity_type == 'caixa' else 'un'}"
            linha[5].text = a.operador or "—"
            linha[6].text = a.motivo or "—"

    nome = f"relatorio-avarias-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M')}.docx"
    return _resposta_docx(doc, nome)


@bp.get("/perdas.docx")
@login_obrigatorio(perfis=["admin"])
def relatorio_perdas():
    """Admin-only — mesma regra de acesso de GET /api/metricas/perdas, já
    que é literalmente a mesma agregação (`calcular_perdas_por_mercado`),
    só que formatada como documento em vez de JSON."""
    dias = int(request.args.get("dias", 30))
    dados = calcular_perdas_por_mercado(g.empresa_id, dias)

    doc = Document()
    _cabecalho(
        doc,
        "Relatório Geral de Perdas",
        f"Últimos {dias} dias — {len(dados['mercados'])} mercado(s) com movimentação de saída",
    )

    if not dados["mercados"]:
        doc.add_paragraph("Nenhuma saída (venda/avaria/vencimento) registrada na janela selecionada.")
    else:
        tabela = _tabela(doc, ["Mercado", "Avaria", "Vencimento", "Venda", "Total perdas", "% perdas"])
        for m in dados["mercados"]:
            linha = tabela.add_row().cells
            linha[0].text = m["mercadoNome"]
            linha[1].text = str(m["avaria"])
            linha[2].text = str(m["vencimento"])
            linha[3].text = str(m["venda"])
            linha[4].text = str(m["totalPerdas"])
            linha[5].text = f"{m['percentual']}%"

    nome = f"relatorio-perdas-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M')}.docx"
    return _resposta_docx(doc, nome)


@bp.get("/giro-registrado.docx")
@login_obrigatorio()
def relatorio_giro_registrado():
    """Giro Registrado ("Qtd. a pedir") de UM mercado + UM ciclo — reflexo
    da aba "Giro Registrado" do Painel de Sugestão. Pedido do usuário
    (2026-10-14): o relatório que existia antes ("Baixar JSON do estoque")
    despejava TODOS os mercados/clientes misturados num arquivo só, difícil
    de ler — este aqui é só o que está na tela daquele mercado naquele
    ciclo. Formatação deliberadamente SEM tabela (um produto por linha) —
    pedido explícito do usuário, pra não quebrar estranho ao abrir no
    Word/celular. Aberto pros dois perfis, mesmo nível de sensibilidade de
    /api/avarias e /api/pedidos: o operador também quer poder baixar o
    giro do mercado em que está trabalhando."""
    mercado_slug = (request.args.get("mercadoId") or "").strip()
    ciclo_id = (request.args.get("cicloId") or "").strip()
    if not mercado_slug or not ciclo_id:
        return jsonify({"erro": "Parâmetros mercadoId e cicloId são obrigatórios."}), 400

    mercado = Mercado.query.filter_by(empresa_id=g.empresa_id, slug=mercado_slug).first()
    nome_mercado = mercado.nome if mercado else mercado_slug

    itens = (
        PedidoItem.query.filter_by(empresa_id=g.empresa_id, mercado_slug=mercado_slug, ciclo_id=ciclo_id)
        .order_by(PedidoItem.atualizado_em.desc())
        .all()
    )

    # Descrição vem do catálogo (Produto) — o PedidoItem só guarda o EAN.
    eans = [i.codigo_ean for i in itens]
    produtos_por_ean = (
        {p.codigo_ean: p for p in Produto.query.filter_by(empresa_id=g.empresa_id).filter(Produto.codigo_ean.in_(eans))}
        if eans
        else {}
    )

    doc = Document()
    doc.add_heading("Giro Registrado", level=1)
    p_mercado = doc.add_paragraph(nome_mercado)
    p_mercado.runs[0].bold = True
    p_mercado.runs[0].font.size = Pt(14)
    agora = datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC")
    meta = doc.add_paragraph(f"Ciclo {ciclo_id}  •  {len(itens)} item(ns)  •  Gerado em {agora}")
    meta.runs[0].font.size = Pt(9)
    doc.add_paragraph()

    rupturas = 0
    if not itens:
        doc.add_paragraph("Nenhum item registrado neste ciclo ainda.")
    else:
        # Ordena por descrição — mais fácil de achar um produto específico
        # numa lista impressa do que pela ordem de edição.
        itens_ordenados = sorted(
            itens, key=lambda i: (produtos_por_ean[i.codigo_ean].descricao if i.codigo_ean in produtos_por_ean else i.codigo_ean)
        )
        for item in itens_ordenados:
            produto = produtos_por_ean.get(item.codigo_ean)
            descricao = produto.descricao if produto else f"EAN {item.codigo_ean}"
            paragrafo = doc.add_paragraph()
            if item.quantidade == 0:
                rupturas += 1
                run = paragrafo.add_run(f"{descricao} — RUPTURA (sem estoque)")
                run.bold = True
            else:
                rotulo = "caixa" if item.quantity_type == "caixa" else "unidade"
                if item.quantidade != 1:
                    rotulo += "s"
                paragrafo.add_run(f"{descricao} — {item.quantidade} {rotulo}")
            run_operador = paragrafo.add_run(f"   (operador: {item.operador or '—'})")
            run_operador.font.size = Pt(9)
            run_operador.italic = True
            paragrafo.paragraph_format.space_after = Pt(6)

        doc.add_paragraph()
        resumo = doc.add_paragraph()
        resumo.add_run(f"Total de itens: {len(itens)}   •   Rupturas: {rupturas}").bold = True

    nome = f"giro-registrado-{mercado_slug}-{ciclo_id}.docx"
    return _resposta_docx(doc, nome)
